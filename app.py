# -*- coding: utf-8 -*-
"""
FORGE — serveur web du questionnaire "Capital récupération" + CRM privé
========================================================================

  • questionnaire public                       ->  GET  /
  • réception d'une soumission (SQLite)        ->  POST /api/submit
  • espace équipe (login mot de passe)         ->  GET  /admin
  • détail d'une personne                      ->  GET  /admin/<id>
  • PDF de toutes les réponses d'une personne  ->  GET  /admin/<id>/pdf
  • ZIP de TOUS les PDF                         ->  GET  /admin/export_pdfs.zip
  • Excel synthèse                              ->  GET  /admin/export.xlsx
  • export CSV brut                             ->  GET  /admin/export.csv

Lancement local :
    pip install -r requirements.txt
    export FORGE_ADMIN_PASSWORD="..."  FORGE_SECRET_KEY="..."
    python app.py            # http://127.0.0.1:8000

Déploiement : voir DEPLOIEMENT.md.
"""

import os
import io
import csv
import json
import zipfile
import sqlite3
import secrets
import datetime
from functools import wraps
from flask import (Flask, request, jsonify, Response, redirect, url_for,
                   session, render_template_string, send_file, abort)

# ----------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("FORGE_DB_PATH", os.path.join(BASE_DIR, "forge_crm.db"))
QUIZ_FILE = os.path.join(BASE_DIR, "forge-quiz-v8.html")

ADMIN_USER = os.environ.get("FORGE_ADMIN_USER", "forge")
ADMIN_PASSWORD = os.environ.get("FORGE_ADMIN_PASSWORD", "change-moi")
SECRET_KEY = os.environ.get("FORGE_SECRET_KEY", secrets.token_hex(32))

app = Flask(__name__)
app.secret_key = SECRET_KEY
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("FORGE_HTTPS", "0") == "1",
    MAX_CONTENT_LENGTH=512 * 1024,
)

# ----------------------------------------------------------------------
# Libellés (source unique = le questionnaire Python, si présent)
# ----------------------------------------------------------------------
Q_LABEL = {}      # qid -> texte de la question
OPT_LABEL = {}    # qid -> {valeur: libellé}
PAGE_OF = {}      # qid -> section
try:
    import forge_questionnaire as FQ
    for pg in FQ.PAGES:
        for q in pg["qs"]:
            Q_LABEL[q["id"]] = q["l"]
            PAGE_OF[q["id"]] = pg["sec"]
            if q.get("o"):
                OPT_LABEL[q["id"]] = {v: t for v, t in q["o"]}
    AXLBL = FQ.AXLBL
except Exception:
    AXLBL = {"sommeil": "Sommeil", "forme": "Forme physique", "recup": "Récupération",
             "stress": "Stress", "confort": "Confort / douleurs", "nutrition": "Nutrition"}

# libellés supplémentaires non-questions
Q_LABEL.setdefault("retraite", "Êtes-vous à la retraite ?")
Q_LABEL.setdefault("enfants3", "Enfant(s) de moins de 3 ans à la maison ?")
OPT_LABEL.setdefault("retraite", {"oui": "Oui", "non": "Non"})
OPT_LABEL.setdefault("enfants3", {"oui": "Oui", "non": "Non"})

FREQ_LABEL = {"f0": "0 / sem", "f12": "1–2 / sem", "f34": "3–4 / sem", "f5": "5+ / sem"}
COMPET_LABEL = {"comp_reg": "Oui, régulièrement", "comp_occ": "Oui, parfois",
                "loisir": "Non (loisir)", "nonsport": "Non (pas de sport)"}
PREF_LABEL = {"passive": "Passive / sensorielle / high-tech", "active": "Active (mouvement)",
              "manuelle": "Manuelle", "mentale": "Mentale", "sociale": "Sociale / conviviale"}

# ----------------------------------------------------------------------
# Base de données
# ----------------------------------------------------------------------
def _resolve_db_path(path):
    try:
        parent = os.path.dirname(path) or "."
        os.makedirs(parent, exist_ok=True)
        t = os.path.join(parent, ".forge_write_test")
        with open(t, "w") as fh:
            fh.write("ok")
        os.remove(t)
        return path
    except Exception:
        fb = os.path.join(BASE_DIR, "forge_crm.db")
        print(f"[FORGE] '{path}' non accessible, bascule sur {fb}")
        return fb


DB_PATH = _resolve_db_path(DB_PATH)


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


COLUMNS = [
    ("created_at", "TEXT"), ("prenom", "TEXT"), ("nom", "TEXT"), ("age", "INTEGER"),
    ("email", "TEXT"), ("code_postal", "TEXT"), ("metier", "TEXT"),
    ("retraite", "INTEGER"), ("enfants_moins3", "INTEGER"),
    ("activite", "TEXT"), ("senior", "INTEGER"),
    ("objectif", "TEXT"), ("objectifs_sec", "TEXT"), ("preference", "TEXT"),
    ("persona", "TEXT"), ("persona_stage", "TEXT"), ("persona_sport", "TEXT"),
    ("persona_douleur", "INTEGER"),
    ("score_global", "INTEGER"), ("meilleur_axe", "TEXT"), ("axe_faible", "TEXT"),
    ("s_sommeil", "REAL"), ("s_forme", "REAL"), ("s_recup", "REAL"),
    ("s_stress", "REAL"), ("s_confort", "REAL"), ("s_nutrition", "REAL"),
    ("besoins", "TEXT"), ("newsletter", "INTEGER"), ("consent_recherche", "INTEGER"),
    ("answers_json", "TEXT"), ("ip", "TEXT"), ("user_agent", "TEXT"),
]


def init_db():
    conn = get_db()
    cols_sql = ",\n".join(f"{n} {t}" for n, t in COLUMNS)
    conn.execute(f"CREATE TABLE IF NOT EXISTS submissions (id INTEGER PRIMARY KEY AUTOINCREMENT,\n{cols_sql})")
    conn.commit()
    # migration douce : ajoute les colonnes manquantes sur une base existante
    existing = {r["name"] for r in conn.execute("PRAGMA table_info(submissions)")}
    for n, t in COLUMNS:
        if n not in existing:
            try:
                conn.execute(f"ALTER TABLE submissions ADD COLUMN {n} {t}")
            except Exception:
                pass
    conn.commit()
    conn.close()


init_db()

# ----------------------------------------------------------------------
# Persona lisible (calculé côté serveur = source de vérité)
# ----------------------------------------------------------------------
def pain_active(ans):
    z = ans.get("d_zone") or []
    return isinstance(z, list) and len(z) > 0 and "none" not in z


def is_sportif(ans):
    return ans.get("compet") in ("comp_reg", "comp_occ") or ans.get("freq") in ("f34", "f5")


def compute_persona(ans, retraite, enfants3):
    if retraite:
        stage = "Retraité"
    elif enfants3:
        stage = "Jeune parent"
    else:
        stage = "Adulte"
    sport = "sportif" if is_sportif(ans) else "non sportif"
    chronic = pain_active(ans) and ans.get("d_anc") == "a3plus"
    label = f"{stage} {sport}" + (" douloureux" if chronic else "")
    return {"label": label, "stage": stage, "sport": sport, "douleur": chronic}


# ----------------------------------------------------------------------
# Auth
# ----------------------------------------------------------------------
def login_required(f):
    @wraps(f)
    def wrapper(*a, **k):
        if not session.get("auth"):
            return redirect(url_for("login", next=request.path))
        return f(*a, **k)
    return wrapper


# ----------------------------------------------------------------------
# Public
# ----------------------------------------------------------------------
@app.route("/")
def quiz():
    try:
        with open(QUIZ_FILE, encoding="utf-8") as fh:
            return Response(fh.read(), mimetype="text/html")
    except FileNotFoundError:
        return Response("forge-quiz-v8.html manquant.", status=500, mimetype="text/plain")


def _int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _flt(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


@app.route("/api/submit", methods=["POST"])
def api_submit():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"ok": False, "error": "payload invalide"}), 400
    ans = data.get("answers") or {}
    sc = data.get("scores") or {}
    age = _int(data.get("age"))
    # Retraité déduit de l'âge (seuil 60 ans) — plus de question dans le questionnaire
    retraite = 1 if (age is not None and age >= 60) else 0
    enfants3 = 1 if (data.get("enfantsMoins3") or ans.get("enfants3") == "oui") else 0
    p = compute_persona(ans, retraite, enfants3)
    row = {
        "created_at": data.get("submittedAt") or datetime.datetime.utcnow().isoformat() + "Z",
        "prenom": (data.get("prenom") or "")[:120],
        "nom": (data.get("nom") or "")[:120],
        "age": age,
        "email": (data.get("email") or "")[:200],
        "code_postal": (data.get("codePostal") or "")[:12],
        "metier": (data.get("metier") or "")[:200],
        "retraite": retraite,
        "enfants_moins3": enfants3,
        "activite": (data.get("activite") or "")[:60],
        "senior": 1 if data.get("senior") else 0,
        "objectif": (data.get("objectifPrioritaire") or "")[:60],
        "objectifs_sec": json.dumps(data.get("objectifsSecondaires") or [], ensure_ascii=False),
        "preference": (data.get("preference") or "")[:60],
        "persona": p["label"],
        "persona_stage": p["stage"],
        "persona_sport": p["sport"],
        "persona_douleur": 1 if p["douleur"] else 0,
        "score_global": _int(data.get("scoreGlobal")),
        "meilleur_axe": (data.get("meilleurAxe") or "")[:40],
        "axe_faible": (data.get("axeFaible") or "")[:40],
        "s_sommeil": _flt(sc.get("sommeil")), "s_forme": _flt(sc.get("forme")),
        "s_recup": _flt(sc.get("recup")), "s_stress": _flt(sc.get("stress")),
        "s_confort": _flt(sc.get("confort")), "s_nutrition": _flt(sc.get("nutrition")),
        "besoins": json.dumps(data.get("besoins") or [], ensure_ascii=False),
        "newsletter": 1 if data.get("newsletter") else 0,
        "consent_recherche": 1 if data.get("consentementRecherche") else 0,
        "answers_json": json.dumps(ans, ensure_ascii=False)[:60000],
        "ip": (request.headers.get("X-Forwarded-For", request.remote_addr) or "")[:80],
        "user_agent": (request.headers.get("User-Agent") or "")[:300],
    }
    cols = ",".join(row.keys())
    ph = ",".join("?" for _ in row)
    conn = get_db()
    cur = conn.execute(f"INSERT INTO submissions ({cols}) VALUES ({ph})", list(row.values()))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return jsonify({"ok": True, "id": new_id})


# ----------------------------------------------------------------------
# Auth routes
# ----------------------------------------------------------------------
@app.route("/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        u = request.form.get("user", "")
        p = request.form.get("password", "")
        if secrets.compare_digest(u, ADMIN_USER) and secrets.compare_digest(p, ADMIN_PASSWORD):
            session["auth"] = True
            session["user"] = u
            return redirect(request.args.get("next") or url_for("admin"))
        error = "Identifiants incorrects."
    return render_template_string(LOGIN_HTML, error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ----------------------------------------------------------------------
# Helpers d'analyse (pour Excel + PDF)
# ----------------------------------------------------------------------
def _answers(r):
    try:
        return json.loads(r["answers_json"] or "{}")
    except Exception:
        return {}


def synth_fields(r):
    """Dérive les colonnes de synthèse demandées à partir d'une ligne."""
    ans = _answers(r)
    sportif = "Oui" if (r["persona_sport"] == "sportif") else "Non"
    volume = FREQ_LABEL.get(ans.get("freq"), "—")
    compet = COMPET_LABEL.get(ans.get("compet"), "—")
    if pain_active(ans):
        douleur = "Oui — chronique" if ans.get("d_anc") == "a3plus" else "Oui"
    else:
        douleur = "Non"
    bstress = ans.get("b_stress")
    stress = "Oui" if (isinstance(bstress, (int, float)) and bstress >= 4) else "Non"
    pref = PREF_LABEL.get(ans.get("pref"), "—")
    faible = AXLBL.get(r["axe_faible"], r["axe_faible"] or "—")
    fort = AXLBL.get(r["meilleur_axe"], r["meilleur_axe"] or "—")
    return {"sportif": sportif, "volume": volume, "compet": compet, "douleur": douleur,
            "stress": stress, "pref": pref, "faible": faible, "fort": fort}


# ----------------------------------------------------------------------
# PDF d'une personne
# ----------------------------------------------------------------------
def build_pdf(r):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm, title=f"Forge #{r['id']}")
    ss = getSampleStyleSheet()
    H = ParagraphStyle("H", parent=ss["Title"], fontSize=20, spaceAfter=2, textColor=colors.black)
    SUB = ParagraphStyle("SUB", parent=ss["Normal"], fontSize=10, textColor=colors.grey, spaceAfter=10)
    SEC = ParagraphStyle("SEC", parent=ss["Heading2"], fontSize=12, spaceBefore=12, spaceAfter=4,
                         textColor=colors.HexColor("#333333"))
    N = ss["Normal"]
    story = []

    story.append(Paragraph("FORGE", H))
    story.append(Paragraph("Bilan capital récupération — fiche individuelle", SUB))

    # Identité
    story.append(Paragraph("Profil", SEC))
    ans = _answers(r)
    ident = [
        ["Prénom", r["prenom"] or "—", "Nom", r["nom"] or "—"],
        ["Âge", str(r["age"] or "—"), "Code postal", r["code_postal"] or "—"],
        ["Email", r["email"] or "—", "Métier", r["metier"] or "—"],
        ["Retraité", "Oui" if r["retraite"] else "Non", "Enfant < 3 ans", "Oui" if r["enfants_moins3"] else "Non"],
        ["Persona", r["persona"] or "—", "Date", (r["created_at"] or "").replace("T", " ")[:16]],
    ]
    t = Table(ident, colWidths=[30 * mm, 55 * mm, 30 * mm, 55 * mm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.grey),
        ("TEXTCOLOR", (2, 0), (2, -1), colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
    ]))
    story.append(t)

    # Scores
    story.append(Paragraph(f"Score global : <b>{r['score_global']}%</b>", SEC))
    srows = [["Dimension", "Score"]]
    for k, lab in AXLBL.items():
        v = r["s_" + k]
        srows.append([lab, f"{round(v)} %" if v is not None else "—"])
    ts = Table(srows, colWidths=[90 * mm, 30 * mm])
    ts.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A1A1A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F3F3")]),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(ts)
    pf = synth_fields(r)
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"Point fort : <b>{pf['fort']}</b> &nbsp;·&nbsp; Point faible : <b>{pf['faible']}</b>", N))

    # Toutes les réponses
    story.append(Paragraph("Réponses au questionnaire", SEC))
    qrows = [["Question", "Réponse"]]
    for qid, val in ans.items():
        qlabel = Q_LABEL.get(qid, qid)
        if isinstance(val, list):
            disp = ", ".join(OPT_LABEL.get(qid, {}).get(v, str(v)) for v in val) or "—"
        elif qid in OPT_LABEL:
            disp = OPT_LABEL[qid].get(val, str(val))
        else:
            disp = str(val)
        qrows.append([Paragraph(qlabel, N), Paragraph(disp, N)])
    qt = Table(qrows, colWidths=[105 * mm, 65 * mm], repeatRows=1)
    qt.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A1A1A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E2E2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(qt)
    story.append(Spacer(1, 8))
    story.append(Paragraph("Bilan d'orientation bien-être — ne remplace pas un avis médical.", SUB))

    doc.build(story)
    buf.seek(0)
    return buf


def _safe_name(r):
    base = f"{(r['nom'] or '').strip()}_{(r['prenom'] or '').strip()}_{r['id']}"
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in base) or f"forge_{r['id']}"


@app.route("/admin/<int:sub_id>/pdf")
@login_required
def admin_pdf(sub_id):
    conn = get_db()
    r = conn.execute("SELECT * FROM submissions WHERE id=?", (sub_id,)).fetchone()
    conn.close()
    if not r:
        abort(404)
    buf = build_pdf(r)
    return send_file(buf, mimetype="application/pdf", as_attachment=True,
                     download_name=_safe_name(r) + ".pdf")


@app.route("/admin/export_pdfs.zip")
@login_required
def export_pdfs_zip():
    conn = get_db()
    rows = conn.execute("SELECT * FROM submissions ORDER BY id").fetchall()
    conn.close()
    mem = io.BytesIO()
    with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as zf:
        for r in rows:
            try:
                zf.writestr(_safe_name(r) + ".pdf", build_pdf(r).read())
            except Exception as e:
                zf.writestr(f"ERREUR_{r['id']}.txt", f"PDF non généré : {e}")
    mem.seek(0)
    fn = "forge_pdfs_" + datetime.date.today().isoformat() + ".zip"
    return send_file(mem, mimetype="application/zip", as_attachment=True, download_name=fn)


@app.route("/admin/export.xlsx")
@login_required
def export_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    conn = get_db()
    rows = conn.execute("SELECT * FROM submissions ORDER BY id DESC").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    headers = ["Prénom", "Nom", "Code postal", "Email",
               "Sportif", "Volume", "Compétition",
               "Douleur", "Stress", "Préférence d'approche",
               "Point faible", "Point fort",
               "Persona", "Score global", "Retraité", "Enfant <3 ans",
               "Newsletter", "Consent. recherche", "Date"]
    ws.append(headers)
    hdr_fill = PatternFill("solid", fgColor="1A1A1A")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(vertical="center")
    for r in rows:
        f = synth_fields(r)
        ws.append([
            r["prenom"], r["nom"], r["code_postal"], r["email"],
            f["sportif"], f["volume"], f["compet"],
            f["douleur"], f["stress"], f["pref"],
            f["faible"], f["fort"],
            r["persona"], r["score_global"],
            "Oui" if r["retraite"] else "Non", "Jeune parent" if r["enfants_moins3"] else "Non",
            "Oui" if r["newsletter"] else "Non", "Oui" if r["consent_recherche"] else "Non",
            (r["created_at"] or "").replace("T", " ")[:16],
        ])
    widths = [12, 14, 11, 26, 8, 11, 18, 16, 8, 26, 16, 16, 26, 11, 9, 12, 10, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    fn = "forge_synthese_" + datetime.date.today().isoformat() + ".xlsx"
    return send_file(mem, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=fn)


@app.route("/admin/export.csv")
@login_required
def export_csv():
    conn = get_db()
    rows = conn.execute("SELECT * FROM submissions ORDER BY id").fetchall()
    conn.close()
    out = io.StringIO()
    if rows:
        w = csv.writer(out)
        w.writerow(rows[0].keys())
        for r in rows:
            w.writerow([r[k] for k in r.keys()])
    data = out.getvalue().encode("utf-8-sig")
    fn = "forge_crm_" + datetime.date.today().isoformat() + ".csv"
    return send_file(io.BytesIO(data), mimetype="text/csv", as_attachment=True, download_name=fn)


# ----------------------------------------------------------------------
# Dashboard
# ----------------------------------------------------------------------
@app.route("/admin")
@login_required
def admin():
    q = (request.args.get("q") or "").strip()
    persona_f = (request.args.get("persona") or "").strip()
    consent_f = request.args.get("consent") or ""
    news_f = request.args.get("news") or ""
    where, params = [], []
    if q:
        where.append("(prenom LIKE ? OR nom LIKE ? OR email LIKE ? OR metier LIKE ? OR persona LIKE ? OR code_postal LIKE ?)")
        params += [f"%{q}%"] * 6
    if persona_f:
        where.append("persona LIKE ?")
        params.append(f"%{persona_f}%")
    if consent_f in ("0", "1"):
        where.append("consent_recherche = ?")
        params.append(int(consent_f))
    if news_f in ("0", "1"):
        where.append("newsletter = ?")
        params.append(int(news_f))
    clause = ("WHERE " + " AND ".join(where)) if where else ""
    conn = get_db()
    rows = conn.execute(f"SELECT * FROM submissions {clause} ORDER BY id DESC LIMIT 500", params).fetchall()
    total = conn.execute("SELECT COUNT(*) c FROM submissions").fetchone()["c"]
    news_count = conn.execute("SELECT COUNT(*) c FROM submissions WHERE newsletter=1").fetchone()["c"]
    rech_count = conn.execute("SELECT COUNT(*) c FROM submissions WHERE consent_recherche=1").fetchone()["c"]
    avg_global = conn.execute("SELECT AVG(score_global) a FROM submissions").fetchone()["a"]
    axis_avgs = conn.execute(
        "SELECT AVG(s_sommeil) so,AVG(s_forme) fo,AVG(s_recup) re,AVG(s_stress) st,AVG(s_confort) co,AVG(s_nutrition) nu FROM submissions").fetchone()
    top_personas = conn.execute(
        "SELECT persona, COUNT(*) c FROM submissions WHERE persona IS NOT NULL AND persona<>'' GROUP BY persona ORDER BY c DESC LIMIT 10").fetchall()
    conn.close()
    stats = {"total": total, "news": news_count, "rech": rech_count,
             "avg_global": round(avg_global) if avg_global else 0,
             "axes": {"Sommeil": axis_avgs["so"], "Forme": axis_avgs["fo"], "Récup.": axis_avgs["re"],
                      "Stress": axis_avgs["st"], "Confort": axis_avgs["co"], "Nutrition": axis_avgs["nu"]},
             "personas": top_personas}
    return render_template_string(ADMIN_HTML, rows=rows, stats=stats, q=q, persona_f=persona_f,
                                  consent_f=consent_f, news_f=news_f, user=session.get("user"))


@app.route("/admin/<int:sub_id>")
@login_required
def admin_detail(sub_id):
    conn = get_db()
    r = conn.execute("SELECT * FROM submissions WHERE id=?", (sub_id,)).fetchone()
    conn.close()
    if not r:
        abort(404)
    ans = _answers(r)
    # réponses lisibles
    nice = []
    for qid, val in ans.items():
        lab = Q_LABEL.get(qid, qid)
        if isinstance(val, list):
            disp = ", ".join(OPT_LABEL.get(qid, {}).get(v, str(v)) for v in val) or "—"
        elif qid in OPT_LABEL:
            disp = OPT_LABEL[qid].get(val, str(val))
        else:
            disp = str(val)
        nice.append((lab, disp))
    besoins = []
    try:
        besoins = json.loads(r["besoins"] or "[]")
    except Exception:
        pass
    scores = {"Sommeil": r["s_sommeil"], "Forme": r["s_forme"], "Récup.": r["s_recup"],
              "Stress": r["s_stress"], "Confort": r["s_confort"], "Nutrition": r["s_nutrition"]}
    return render_template_string(DETAIL_HTML, r=r, nice=nice, besoins=besoins,
                                  scores=scores, synth=synth_fields(r))


@app.route("/admin/<int:sub_id>/delete", methods=["POST"])
@login_required
def admin_delete(sub_id):
    conn = get_db()
    conn.execute("DELETE FROM submissions WHERE id=?", (sub_id,))
    conn.commit()
    conn.close()
    return redirect(url_for("admin"))


@app.route("/healthz")
def healthz():
    return "ok"


# ======================================================================
# Templates
# ======================================================================
BASE_CSS = """
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:-apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;background:#101010;color:#F5F5F5;line-height:1.5}
a{color:#F5F5F5}
.wrap{max-width:1240px;margin:0 auto;padding:24px}
.top{display:flex;align-items:center;justify-content:space-between;border-bottom:1px solid rgba(255,255,255,.14);padding-bottom:16px;margin-bottom:22px}
.brand{font-weight:700;letter-spacing:.28em;font-size:18px}
.muted{color:#9A9A9A}
.btn{display:inline-block;background:#fff;color:#101010;border:none;border-radius:9px;padding:9px 16px;font-weight:600;cursor:pointer;text-decoration:none;font-size:14px}
.btn.ghost{background:transparent;color:#F5F5F5;border:1.5px solid rgba(255,255,255,.3)}
.btn.sm{padding:6px 11px;font-size:13px}
.btn.danger{border-color:#C98B6B;color:#E7B597;background:transparent}
.cards{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:20px}
.card{background:#181818;border:1px solid rgba(255,255,255,.12);border-radius:12px;padding:16px 18px;flex:1;min-width:150px}
.card .n{font-size:30px;font-weight:300}
.card .l{font-size:12px;color:#9A9A9A;text-transform:uppercase;letter-spacing:.08em}
.bar{display:flex;align-items:center;gap:8px;margin:4px 0;font-size:13px}
.bar .lab{width:78px;color:#C9C9C9}
.bar .track{flex:1;height:8px;background:rgba(255,255,255,.08);border-radius:5px;overflow:hidden}
.bar .fill{height:100%;background:#fff;border-radius:5px}
.bar .v{width:40px;text-align:right;color:#C9C9C9}
table{width:100%;border-collapse:collapse;font-size:13.5px;margin-top:8px}
th,td{text-align:left;padding:9px 10px;border-bottom:1px solid rgba(255,255,255,.08);white-space:nowrap}
th{color:#9A9A9A;font-weight:600;font-size:12px;text-transform:uppercase;letter-spacing:.04em}
tr:hover td{background:rgba(255,255,255,.03)}
.pill{display:inline-block;padding:2px 9px;border:1px solid rgba(255,255,255,.22);border-radius:20px;font-size:12px}
.pill.stage{background:rgba(255,255,255,.1)}
.pill.sport{border-color:rgba(255,255,255,.4)}
.pill.pain{background:rgba(224,166,106,.18);border-color:rgba(224,166,106,.6);color:#E7B597}
form.filters{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-bottom:16px}
input,select{background:#181818;border:1.5px solid rgba(255,255,255,.18);color:#F5F5F5;border-radius:9px;padding:9px 12px;font-size:14px;font-family:inherit}
input::placeholder{color:#777}
.kv{display:grid;grid-template-columns:210px 1fr;gap:8px 16px;font-size:14px;margin-top:10px}
.kv div:nth-child(odd){color:#9A9A9A}
.sect{margin:26px 0 10px;font-size:12px;letter-spacing:.14em;text-transform:uppercase;color:#9A9A9A;border-bottom:1px solid rgba(255,255,255,.1);padding-bottom:6px}
.toolbar{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0 2px}
"""

LOGIN_HTML = """<!doctype html><html lang=fr><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1"><title>Forge — Espace équipe</title>
<style>""" + BASE_CSS + """
.login{max-width:360px;margin:12vh auto;background:#181818;border:1px solid rgba(255,255,255,.12);border-radius:14px;padding:30px}
.login h1{font-weight:300;font-size:22px;margin-bottom:4px}.login .brand{display:block;text-align:center;margin-bottom:18px}
.login input{width:100%;margin-bottom:12px}.err{color:#E7B597;font-size:13px;margin-bottom:10px}
</style></head><body>
<div class=login><span class=brand>F O R G E</span><h1>Espace équipe</h1>
<p class=muted style="font-size:13px;margin-bottom:18px">Accès réservé — données CRM</p>
{% if error %}<div class=err>{{error}}</div>{% endif %}
<form method=post>
<input name=user placeholder="Identifiant" autocomplete=username autofocus>
<input name=password type=password placeholder="Mot de passe" autocomplete=current-password>
<button class=btn style="width:100%" type=submit>Se connecter</button></form></div></body></html>"""

ADMIN_HTML = """<!doctype html><html lang=fr><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1"><title>Forge — CRM</title>
<style>""" + BASE_CSS + """</style></head><body><div class=wrap>
<div class=top><span class=brand>F O R G E · CRM</span>
<span class=muted>{{user}} · <a href="{{url_for('logout')}}">déconnexion</a></span></div>

<div class=cards>
<div class=card><div class=n>{{stats.total}}</div><div class=l>Soumissions</div></div>
<div class=card><div class=n>{{stats.avg_global}}%</div><div class=l>Score global moyen</div></div>
<div class=card><div class=n>{{stats.news}}</div><div class=l>Newsletter OK</div></div>
<div class=card><div class=n>{{stats.rech}}</div><div class=l>Consentement recherche</div></div>
</div>

<div class=cards>
<div class=card style="min-width:320px"><div class=l style="margin-bottom:8px">Moyenne par axe</div>
{% for lab,val in stats.axes.items() %}
<div class=bar><span class=lab>{{lab}}</span><span class=track><span class=fill style="width:{{(val or 0)|round|int}}%"></span></span><span class=v>{{(val or 0)|round|int}}%</span></div>
{% endfor %}</div>
<div class=card style="min-width:320px"><div class=l style="margin-bottom:8px">Personas les plus fréquents</div>
{% for p in stats.personas %}<div class=bar><span class=lab style="width:auto;flex:1">{{p['persona']}}</span><span class=v>{{p['c']}}</span></div>
{% else %}<div class=muted style="font-size:13px">Aucune donnée.</div>{% endfor %}</div>
</div>

<div class=toolbar>
<a class=btn href="{{url_for('export_xlsx')}}">⭳ Excel synthèse</a>
<a class="btn ghost" href="{{url_for('export_pdfs_zip')}}">⭳ Tous les PDF (zip)</a>
<a class="btn ghost" href="{{url_for('export_csv')}}">⭳ CSV brut</a>
</div>

<form class=filters method=get>
<input name=q value="{{q}}" placeholder="Rechercher (nom, email, métier, CP, persona)">
<input name=persona value="{{persona_f}}" placeholder="Persona contient…">
<select name=news>
<option value="" {{'selected' if news_f=='' else ''}}>Newsletter : tous</option>
<option value="1" {{'selected' if news_f=='1' else ''}}>Newsletter : oui</option>
<option value="0" {{'selected' if news_f=='0' else ''}}>Newsletter : non</option></select>
<select name=consent>
<option value="" {{'selected' if consent_f=='' else ''}}>Recherche : tous</option>
<option value="1" {{'selected' if consent_f=='1' else ''}}>Recherche : oui</option>
<option value="0" {{'selected' if consent_f=='0' else ''}}>Recherche : non</option></select>
<button class=btn type=submit>Filtrer</button>
<a class="btn ghost" href="{{url_for('admin')}}">Réinitialiser</a>
</form>

<table>
<tr><th>#</th><th>Date</th><th>Nom</th><th>Prénom</th><th>CP</th><th>Persona</th>
<th>Global</th><th>Faible</th><th>Email</th><th>N</th><th>R</th><th></th></tr>
{% for r in rows %}
<tr>
<td>{{r['id']}}</td>
<td class=muted>{{r['created_at'][:16].replace('T',' ') if r['created_at'] else ''}}</td>
<td>{{r['nom']}}</td><td>{{r['prenom']}}</td><td class=muted>{{r['code_postal']}}</td>
<td>
  <span class="pill stage">{{r['persona_stage']}}</span>
  <span class="pill sport">{{r['persona_sport']}}</span>
  {% if r['persona_douleur'] %}<span class="pill pain">douloureux</span>{% endif %}
</td>
<td style="font-weight:700">{{r['score_global'] if r['score_global'] is not none else ''}}%</td>
<td class=muted>{{r['axe_faible']}}</td>
<td class=muted>{{r['email']}}</td>
<td>{{'✓' if r['newsletter'] else ''}}</td>
<td>{{'✓' if r['consent_recherche'] else ''}}</td>
<td style="white-space:nowrap">
  <a class="btn ghost sm" href="{{url_for('admin_detail', sub_id=r['id'])}}">Voir</a>
  <a class="btn sm" href="{{url_for('admin_pdf', sub_id=r['id'])}}">PDF</a>
</td>
</tr>
{% else %}<tr><td colspan=12 class=muted style="padding:24px">Aucune soumission pour l'instant.</td></tr>{% endfor %}
</table>
<p class=muted style="font-size:12px;margin-top:14px">500 dernières lignes affichées · les exports portent sur toute la base.</p>
</div></body></html>"""

DETAIL_HTML = """<!doctype html><html lang=fr><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1"><title>Forge — #{{r['id']}}</title>
<style>""" + BASE_CSS + """</style></head><body><div class=wrap>
<div class=top><span class=brand>F O R G E · fiche #{{r['id']}}</span>
<span><a class="btn" href="{{url_for('admin_pdf', sub_id=r['id'])}}">⭳ PDF complet</a>
&nbsp;<a class="btn ghost" href="{{url_for('admin')}}">← Retour</a></span></div>

<div class=cards>
<div class=card style="flex:2;min-width:320px"><div class=l>Profil</div>
<div class=kv>
<div>Prénom</div><div>{{r['prenom']}}</div>
<div>Nom</div><div>{{r['nom']}}</div>
<div>Âge</div><div>{{r['age']}}</div>
<div>Code postal</div><div>{{r['code_postal']}}</div>
<div>Email</div><div>{{r['email']}}</div>
<div>Métier</div><div>{{r['metier'] or '—'}}</div>
<div>Retraité</div><div>{{'Oui' if r['retraite'] else 'Non'}}</div>
<div>Enfant &lt; 3 ans</div><div>{{'Oui' if r['enfants_moins3'] else 'Non'}}</div>
<div>Persona</div><div>
  <span class="pill stage">{{r['persona_stage']}}</span>
  <span class="pill sport">{{r['persona_sport']}}</span>
  {% if r['persona_douleur'] %}<span class="pill pain">douloureux</span>{% endif %}</div>
<div>Sportif</div><div>{{synth.sportif}} · {{synth.volume}} · compétition : {{synth.compet}}</div>
<div>Douleur</div><div>{{synth.douleur}}</div>
<div>Stress</div><div>{{synth.stress}}</div>
<div>Préférence</div><div>{{synth.pref}}</div>
<div>Newsletter</div><div>{{'Oui' if r['newsletter'] else 'Non'}}</div>
<div>Consent. recherche</div><div>{{'Oui' if r['consent_recherche'] else 'Non'}}</div>
<div>Date</div><div>{{(r['created_at'] or '').replace('T',' ')}}</div>
</div></div>
<div class=card style="flex:1;min-width:280px"><div class=l>Score global : <b style="color:#fff">{{r['score_global']}}%</b></div>
<div style="margin-top:10px">
{% for lab,val in scores.items() %}
<div class=bar><span class=lab>{{lab}}</span><span class=track><span class=fill style="width:{{(val or 0)|round|int}}%"></span></span><span class=v>{{(val or 0)|round|int}}%</span></div>
{% endfor %}</div>
<p class=muted style="font-size:13px;margin-top:10px">Point fort : <b style="color:#fff">{{synth.fort}}</b><br>Point faible : <b style="color:#fff">{{synth.faible}}</b></p>
</div></div>

<div class=sect>Besoins identifiés (coulisse CRM)</div>
<div>{% for b in besoins %}<span class=pill style="margin:3px">{{b}}</span>{% else %}<span class=muted>—</span>{% endfor %}</div>

<div class=sect>Toutes les réponses</div>
<div class=kv>{% for lab,val in nice %}<div>{{lab}}</div><div>{{val}}</div>{% endfor %}</div>

<div class=sect>Zone RGPD</div>
<form method=post action="{{url_for('admin_delete', sub_id=r['id'])}}" onsubmit="return confirm('Supprimer définitivement cette fiche ?');">
<button class="btn danger sm" type=submit>Supprimer (droit à l'effacement)</button></form>
</div></body></html>"""


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port, debug=False)
